import json
import io
from datetime import datetime, timedelta
from typing import List, Optional
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import func
from database import get_db
from models import (
    DailyReport, OreBatch, MiningEquipment, EquipmentData,
    WorkOrder, Alert, PersonnelCheck, TransportTask, TransportStatus,
    User, UserRole, OreBatchStatus
)
from schemas import DailyReportResponse, AlertResponse
from routers.auth import get_current_user
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

router = APIRouter(prefix="/reports", tags=["运营日报与导出"])


def _generate_daily_report(db: Session, report_date: datetime) -> DailyReport:
    start_of_day = report_date.replace(hour=0, minute=0, second=0, microsecond=0)
    end_of_day = start_of_day + timedelta(days=1)

    area_output = {}
    batches = db.query(OreBatch).filter(
        OreBatch.sampled_at >= start_of_day,
        OreBatch.sampled_at < end_of_day
    ).all()
    total_output = 0.0
    for b in batches:
        area = b.mining_area or "未知区域"
        if area not in area_output:
            area_output[area] = {"output": 0.0, "batch_count": 0, "avg_grade": 0.0, "qualified_count": 0}
        area_output[area]["output"] += b.weight or 0
        area_output[area]["batch_count"] += 1
        if b.grade:
            area_output[area]["avg_grade"] += b.grade
        if b.status in [OreBatchStatus.QUALIFIED, OreBatchStatus.BLENDED]:
            area_output[area]["qualified_count"] += 1
        total_output += b.weight or 0

    for area in area_output:
        if area_output[area]["batch_count"] > 0:
            area_output[area]["avg_grade"] = round(
                area_output[area]["avg_grade"] / area_output[area]["batch_count"], 2
            )

    total_equipments = db.query(MiningEquipment).count()
    abnormal_data_count = db.query(EquipmentData).filter(
        EquipmentData.collected_at >= start_of_day,
        EquipmentData.collected_at < end_of_day,
        EquipmentData.is_abnormal == True
    ).count()
    total_data_count = db.query(EquipmentData).filter(
        EquipmentData.collected_at >= start_of_day,
        EquipmentData.collected_at < end_of_day
    ).count()

    if total_data_count > 0:
        equipment_utilization = round(
            max(0.0, 100.0 - (abnormal_data_count / total_data_count * 100.0)), 2
        )
    else:
        equipment_utilization = 100.0 if total_equipments > 0 else 0.0

    safety_event_count = db.query(Alert).filter(
        Alert.created_at >= start_of_day,
        Alert.created_at < end_of_day
    ).count()
    safety_event_count += db.query(PersonnelCheck).filter(
        PersonnelCheck.check_time >= start_of_day,
        PersonnelCheck.check_time < end_of_day,
        PersonnelCheck.overall_passed == False
    ).count()

    work_order_count = db.query(WorkOrder).filter(
        WorkOrder.created_at >= start_of_day,
        WorkOrder.created_at < end_of_day
    ).count()
    transport_count = db.query(TransportTask).filter(
        TransportTask.created_at >= start_of_day,
        TransportTask.created_at < end_of_day,
        TransportTask.status == TransportStatus.COMPLETED
    ).count()

    area_output["_summary"] = {
        "total_equipments": total_equipments,
        "abnormal_data_count": abnormal_data_count,
        "work_order_count": work_order_count,
        "transport_count": transport_count
    }

    report = DailyReport(
        report_date=start_of_day,
        mining_area_data=json.dumps(area_output, ensure_ascii=False),
        total_output=round(total_output, 2),
        equipment_utilization=equipment_utilization,
        safety_event_count=safety_event_count
    )
    return report


@router.post("/daily/generate", response_model=DailyReportResponse)
def generate_daily_report(
    date: Optional[str] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    if date:
        try:
            report_date = datetime.strptime(date, "%Y-%m-%d")
        except ValueError:
            raise HTTPException(status_code=400, detail="日期格式错误，应为YYYY-MM-DD")
    else:
        report_date = datetime.utcnow() - timedelta(days=1)

    start_of_day = report_date.replace(hour=0, minute=0, second=0, microsecond=0)
    existing = db.query(DailyReport).filter(DailyReport.report_date == start_of_day).first()
    if existing:
        db.delete(existing)
        db.commit()

    report = _generate_daily_report(db, report_date)
    db.add(report)
    db.commit()
    db.refresh(report)
    return report


@router.get("/daily", response_model=List[DailyReportResponse])
def list_daily_reports(
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    query = db.query(DailyReport)
    if start_date:
        try:
            sd = datetime.strptime(start_date, "%Y-%m-%d")
            query = query.filter(DailyReport.report_date >= sd)
        except ValueError:
            raise HTTPException(status_code=400, detail="开始日期格式错误")
    if end_date:
        try:
            ed = datetime.strptime(end_date, "%Y-%m-%d") + timedelta(days=1)
            query = query.filter(DailyReport.report_date < ed)
        except ValueError:
            raise HTTPException(status_code=400, detail="结束日期格式错误")
    return query.order_by(DailyReport.report_date.desc()).all()


@router.get("/daily/{report_id}", response_model=DailyReportResponse)
def get_daily_report(
    report_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    report = db.query(DailyReport).filter(DailyReport.id == report_id).first()
    if not report:
        raise HTTPException(status_code=404, detail="日报不存在")
    return report


@router.get("/daily/{report_id}/detail")
def get_daily_report_detail(
    report_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    report = db.query(DailyReport).filter(DailyReport.id == report_id).first()
    if not report:
        raise HTTPException(status_code=404, detail="日报不存在")
    area_data = json.loads(report.mining_area_data) if report.mining_area_data else {}
    return {
        "report_id": report.id,
        "report_date": report.report_date.isoformat(),
        "total_output": report.total_output,
        "equipment_utilization": report.equipment_utilization,
        "safety_event_count": report.safety_event_count,
        "mining_areas": {k: v for k, v in area_data.items() if k != "_summary"},
        "summary": area_data.get("_summary", {})
    }


@router.get("/export/daily")
def export_daily_report_excel(
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    area: Optional[str] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    query = db.query(DailyReport)
    if start_date:
        try:
            sd = datetime.strptime(start_date, "%Y-%m-%d")
            query = query.filter(DailyReport.report_date >= sd)
        except ValueError:
            raise HTTPException(status_code=400, detail="开始日期格式错误")
    if end_date:
        try:
            ed = datetime.strptime(end_date, "%Y-%m-%d") + timedelta(days=1)
            query = query.filter(DailyReport.report_date < ed)
        except ValueError:
            raise HTTPException(status_code=400, detail="结束日期格式错误")

    reports = query.order_by(DailyReport.report_date.asc()).all()
    if not reports:
        raise HTTPException(status_code=404, detail="无报表数据")

    wb = Workbook()
    ws = wb.active
    ws.title = "矿山运营日报"

    header_font = Font(bold=True, color="FFFFFF", size=12)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")

    headers = ["日期", "总产量(吨)", "设备利用率(%)", "安全事件数", "各采区产量详情"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    for row, report in enumerate(reports, 2):
        area_data = json.loads(report.mining_area_data) if report.mining_area_data else {}
        area_details = []
        for k, v in area_data.items():
            if k != "_summary" and (not area or area == k):
                area_details.append(f"{k}: {v.get('output', 0)}吨")

        ws.cell(row=row, column=1, value=report.report_date.strftime("%Y-%m-%d"))
        ws.cell(row=row, column=2, value=report.total_output)
        ws.cell(row=row, column=3, value=report.equipment_utilization)
        ws.cell(row=row, column=4, value=report.safety_event_count)
        ws.cell(row=row, column=5, value="; ".join(area_details))

    ws_detail = wb.create_sheet("采区明细")
    detail_headers = ["日期", "采区", "产量(吨)", "批次数", "平均品位(%)", "合格批次数"]
    for col, h in enumerate(detail_headers, 1):
        cell = ws_detail.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    row = 2
    for report in reports:
        area_data = json.loads(report.mining_area_data) if report.mining_area_data else {}
        date_str = report.report_date.strftime("%Y-%m-%d")
        for k, v in area_data.items():
            if k != "_summary" and (not area or area == k):
                ws_detail.cell(row=row, column=1, value=date_str)
                ws_detail.cell(row=row, column=2, value=k)
                ws_detail.cell(row=row, column=3, value=v.get("output", 0))
                ws_detail.cell(row=row, column=4, value=v.get("batch_count", 0))
                ws_detail.cell(row=row, column=5, value=v.get("avg_grade", 0))
                ws_detail.cell(row=row, column=6, value=v.get("qualified_count", 0))
                row += 1

    for col in range(1, 6):
        ws.column_dimensions[chr(64 + col)].width = 20
    ws.column_dimensions["E"].width = 60
    for col in range(1, 7):
        ws_detail.column_dimensions[chr(64 + col)].width = 18

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"矿山运营日报_{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@router.get("/alerts", response_model=List[AlertResponse])
def list_alerts(
    level: Optional[str] = None,
    resolved: Optional[bool] = None,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    limit: int = 100,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    query = db.query(Alert)
    if level:
        query = query.filter(Alert.level == level)
    if resolved is not None:
        query = query.filter(Alert.is_resolved == resolved)
    if start_date:
        sd = datetime.strptime(start_date, "%Y-%m-%d")
        query = query.filter(Alert.created_at >= sd)
    if end_date:
        ed = datetime.strptime(end_date, "%Y-%m-%d") + timedelta(days=1)
        query = query.filter(Alert.created_at < ed)
    return query.order_by(Alert.created_at.desc()).limit(limit).all()


@router.put("/alerts/{alert_id}/resolve", response_model=AlertResponse)
def resolve_alert(
    alert_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    alert = db.query(Alert).filter(Alert.id == alert_id).first()
    if not alert:
        raise HTTPException(status_code=404, detail="告警不存在")
    alert.is_resolved = True
    alert.resolved_at = datetime.utcnow()
    db.commit()
    db.refresh(alert)
    return alert
